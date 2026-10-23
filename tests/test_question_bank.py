import base64
from io import BytesIO
import json
from pathlib import Path
import threading
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import Workbook
import pytest

from smallai.data import is_validation, iter_jsonl
from smallai.public_data import PreparationCancelled
from smallai.question_bank import prepare_question_bank, preview_question_bank, question_package
from smallai.tokenizer import sha256_file
from smallai.tools import compact_json
from smallai.ui import BusyError, TrainingManager, create_server
from test_ui import wait_until


def encoded(content):
    return base64.b64encode(content if isinstance(content, bytes) else content.encode("utf-8")).decode()


def bank(count=80):
    return [{"题目": f"计算 {index} 加 2 等于多少？", "答案": str(index + 2),
             "解析": f"{index} + 2 = {index + 2}"} for index in range(count)]


def preview(root, rows=None):
    return preview_question_bank(root, "我的题库.json", encoded(json.dumps(rows or bank(), ensure_ascii=False)))


@pytest.mark.parametrize("extension", ["csv", "tsv", "json", "jsonl", "xlsx"])
def test_question_bank_formats_preserve_choices_and_explanation(tmp_path, extension):
    headers = ["题目", "A", "B", "答案", "解析"]
    values = ["一加一等于多少？", "二", "三", "A", "一加一等于二。"]
    if extension == "xlsx":
        workbook = Workbook()
        workbook.active.append(headers)
        workbook.active.append(values)
        stream = BytesIO()
        workbook.save(stream)
        content = stream.getvalue()
    elif extension in ("csv", "tsv"):
        separator = "," if extension == "csv" else "\t"
        content = separator.join(headers) + "\n" + separator.join(values)
    else:
        record = dict(zip(headers, values))
        content = json.dumps([record] if extension == "json" else record, ensure_ascii=False)
    result = preview_question_bank(tmp_path, "bank." + extension, encoded(content))
    assert result["accepted"] == 1
    assert result["samples"][0]["messages"] == [
        {"role": "user", "content": "一加一等于多少？\n\nA. 二\nB. 三"},
        {"role": "assistant", "content": "A\n\n解析：一加一等于二。"}]


def test_aliases_numeric_zero_and_gb18030_csv(tmp_path):
    result = preview_question_bank(tmp_path, "bank.csv", encoded("question,answer,选项A,选项B\n一减一,0,零,一".encode("gb18030")))
    assert result["samples"][0]["messages"][1]["content"] == "0"
    assert "A. 零" in result["samples"][0]["messages"][0]["content"]


def test_invalid_rows_conflicts_and_duplicates_are_reported(tmp_path):
    lines = [compact_json(bank(1)[0]), compact_json(bank(1)[0]), "broken", "[]",
             compact_json({"question": "only question"}),
             compact_json({**bank(1)[0], "答案": "999"}), compact_json(bank(2)[1])]
    result = preview_question_bank(tmp_path, "bank.jsonl", encoded("\n".join(lines)))
    assert (result["total"], result["accepted"], result["duplicates"], result["invalid"]) == (7, 2, 1, 4)
    assert [value["row"] for value in result["errors"]] == [3, 4, 5, 6]
    assert "答案不一致" in result["errors"][-1]["message"]


def test_excel_formula_is_skipped_and_not_executed(tmp_path):
    workbook = Workbook()
    workbook.active.append(["题目", "答案"])
    workbook.active.append(["有效题目", "有效答案"])
    workbook.active.append(["公式题目", "=1+2"])
    workbook.active.append(["数学符号题目", "= 表示相等"])
    workbook.active.cell(4, 2).data_type = "s"
    stream = BytesIO()
    workbook.save(stream)
    result = preview_question_bank(tmp_path, "bank.xlsx", encoded(stream.getvalue()))
    assert result["accepted"] == 2 and result["invalid"] == 1
    assert "公式" in result["errors"][0]["message"]


@pytest.mark.parametrize("filename,content", [("x.txt", "text"), ("x.json", "not json"),
                                             ("x.csv", "题目,题目\n重复,列名"), ("x.json", "[]")])
def test_bad_files_are_rejected_without_registering_datasets(tmp_path, filename, content):
    with pytest.raises(ValueError):
        preview_question_bank(tmp_path, filename, encoded(content))
    assert not (tmp_path / "data/prepared").exists()


def test_new_tokenizer_keeps_validation_questions_out_of_pretraining_and_vocabulary(tmp_path):
    uploaded = preview(tmp_path)
    result = prepare_question_bank(tmp_path, uploaded["id"])
    assert result["accepted"] == 80
    for split in ("train", "val"):
        assert result["prepared"]["sft"][split]["documents"] == result["prepared"]["pretrain"][split]["documents"]
        assert result["prepared"]["sft"][split]["documents"] > 0
    package = question_package(tmp_path, uploaded["id"])
    training_text = (package / "tokenizer_corpus.jsonl").read_text(encoding="utf-8")
    for record in iter_jsonl([package / "records.jsonl"]):
        identity = compact_json({"messages": record["messages"], "tools": []})
        if is_validation(identity):
            assert record["messages"][0]["content"] not in training_text
    with pytest.raises(ValueError, match="准备过"):
        prepare_question_bank(tmp_path, uploaded["id"])


def test_existing_tokenizer_is_reused_and_long_questions_reported(ui_workspace):
    token_path = ui_workspace / "data/tokenizer.json"
    before = sha256_file(token_path)
    uploaded = preview(ui_workspace, bank() + [{"question": "很长的题目" * 1000, "answer": "答案"}])
    result = prepare_question_bank(ui_workspace, uploaded["id"], tokenizer_path=token_path)
    assert sha256_file(token_path) == before
    assert result["accepted"] == 80 and result["too_long_skipped"] == 1
    meta = json.loads((ui_workspace / result["outputs"]["sft"] / "manifest.json").read_text(encoding="utf-8"))
    assert meta["source_type"] == "question_bank" and meta["tokenizer_sha256"] == before


def test_messages_format_and_import_path_and_cancellation(tmp_path):
    records = [{"messages": [{"role": "user", "content": "你好吗？"}, {"role": "assistant", "content": "我很好。"}]}]
    uploaded = preview(tmp_path, records)
    assert uploaded["samples"] == [{**records[0], "tools": []}]
    for identifier in ("../outside", "questions_20261023_000000_abcdef/../../other", None):
        with pytest.raises(ValueError):
            question_package(tmp_path, identifier)
    cancel = threading.Event()
    cancel.set()
    with pytest.raises(PreparationCancelled):
        prepare_question_bank(tmp_path, uploaded["id"], cancel=cancel)
    assert not (tmp_path / "data/prepared").exists()


def test_imported_bank_actually_pretrains_then_learns_answers(ui_workspace):
    manager = TrainingManager(ui_workspace)
    try:
        uploaded = manager.preview_questions({"filename": "试题.json", "content": encoded(json.dumps(bank(), ensure_ascii=False))})
        manager.prepare_questions({"import_id": uploaded["id"], "tokenizer_path": "data/tokenizer.json"})
        with pytest.raises(BusyError):
            manager.prepare_public({})
        wait_until(lambda: manager.state()["data_job"]["status"] not in ("running", "cancelling"))
        job = manager.state()["data_job"]
        assert job["status"] == "completed", job
        output = job["result"]["outputs"]
        assert manager.bootstrap()["datasets"][0]["source_type"] == "question_bank"
        base = manager.start({"data": output["pretrain"], "model": "smoke", "steps": 2, "block_size": 64, "threads": 2})
        wait_until(lambda: manager.state()["active_id"] is None)
        assert manager.state()["run"]["status"] == "completed"
        manager.start({"stage": "sft", "data": output["sft"], "model": "smoke", "steps": 1, "threads": 2,
                       "init_from": base["output"] + "/last.pt"})
        wait_until(lambda: manager.state()["active_id"] is None)
        assert manager.state()["run"]["status"] == "completed", manager.state()
        assert manager.state()["run"]["progress"] == 100
    finally:
        manager.close()


def test_upload_api_uses_csrf_and_bounded_requests(ui_workspace):
    server = create_server(ui_workspace, port=0)
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    address = f"http://127.0.0.1:{server.server_address[1]}"
    try:
        with urlopen(address + "/api/bootstrap") as response:
            bootstrap = json.load(response)
        body = json.dumps({"filename": "题库.json", "content": encoded(json.dumps(bank()))}).encode()
        headers = {"Content-Type": "application/json"}
        with pytest.raises(HTTPError) as rejected:
            urlopen(Request(address + "/api/questions/preview", body, headers))
        assert rejected.value.code == 403
        headers["X-SmallAI-Token"] = bootstrap["csrf_token"]
        with urlopen(Request(address + "/api/questions/preview", body, headers)) as response:
            result = json.load(response)
        assert result["accepted"] == 80
        with urlopen(address + "/api/questions/template.csv") as response:
            assert response.read().decode("utf-8-sig").startswith("题目,A,B")
        huge = Request(address + "/api/questions/preview", b"{}", {**headers, "Content-Length": "20000000"})
        with pytest.raises(HTTPError) as rejected:
            urlopen(huge)
        assert rejected.value.code == 400
    finally:
        server.shutdown()
        server.manager.close()
        server.server_close()
        worker.join(3)
