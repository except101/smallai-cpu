"""Read local question banks and prepare compatible pretraining/SFT datasets."""

import base64
import binascii
import csv
from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
import secrets
import tempfile
import unicodedata
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from .data import is_validation, iter_jsonl, prepare_dataset
from .format import encode_conversation, validate_sft_record
from .public_data import PreparationCancelled
from .tokenizer import LocalTokenizer, sha256_file, train_tokenizer
from .tools import compact_json


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50000
EXTENSIONS = {".csv", ".tsv", ".json", ".jsonl", ".xlsx"}
ALIASES = {
    "question": ("question", "prompt", "instruction", "题目", "题干", "问题"),
    "answer": ("answer", "response", "output", "答案", "正确答案", "参考答案"),
    "explanation": ("explanation", "analysis", "rationale", "解析", "答案解析"),
    "context": ("context", "input", "材料", "上下文"),
    "options": ("options", "choices", "选项"),
}
TEMPLATE = "题目,A,B,C,D,答案,解析\r\n中国的首都是哪里？,北京,上海,广州,深圳,A,北京是中国的首都。\r\n"


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        text = unicodedata.normalize("NFC", str(value).strip())
        if len(text) > 100000:
            raise ValueError("单个字段超过 10 万字符，请拆分题目")
        return text
    raise ValueError("题目、答案和解析应为文字或数字")


def _normalize(row):
    if not isinstance(row, dict):
        raise ValueError("每道题应为一个对象")
    if "messages" in row:
        record = {"messages": row["messages"], "tools": row.get("tools", [])}
        validate_sft_record(record)
        for message in record["messages"]:
            _text(message.get("content"))
        return record
    fields = {str(key).strip().lower(): value for key, value in row.items() if key is not None}

    def field(name):
        return next((fields[key] for key in ALIASES[name] if key in fields and fields[key] not in (None, "")), "")

    question, answer = _text(field("question")), _text(field("answer"))
    if not question or not answer:
        raise ValueError("缺少题目或答案；列名可用「题目 / 答案」或 question / answer")
    context, explanation = _text(field("context")), _text(field("explanation"))
    options = field("options")
    if not options:
        options = {letter: fields.get(letter.lower(), fields.get("选项" + letter.lower(), "")) for letter in "ABCDEFGH"}
    if isinstance(options, dict):
        option_text = "\n".join(f"{key}. {_text(value)}" for key, value in options.items() if value not in (None, ""))
    elif isinstance(options, list):
        option_text = "\n".join(f"{chr(65 + index)}. {_text(value)}" for index, value in enumerate(options))
    else:
        option_text = _text(options)
    prompt = "\n\n".join(part for part in (context, question, option_text) if part)
    response = answer + ("\n\n解析：" + explanation if explanation else "")
    record = {"messages": [{"role": "user", "content": prompt}, {"role": "assistant", "content": response}]}
    validate_sft_record(record)
    return record


def _rows(content, extension):
    try:
        yield from _read_rows(content, extension)
    except csv.Error as error:
        raise ValueError("表格字段过长或格式有误，请检查分隔符和引号") from error
    except RecursionError as error:
        raise ValueError("JSON 嵌套过深，请使用简单的题目数组") from error


def _read_rows(content, extension):
    if extension == ".xlsx":
        try:
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException
        except ImportError as error:
            raise ValueError("读取 Excel 需要 openpyxl，请运行 python -m pip install -e . 或改用 CSV") from error
        try:
            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("Excel 解压后超过 100 MiB，请拆分题库")
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
            try:
                sheet = workbook.active
                # Ignore inaccurate worksheet dimensions supplied by some exporters.
                sheet.reset_dimensions()
                rows = sheet.iter_rows()
                header = [cell.value for cell in next(rows, ())]
                header = _headers(header)
                for number, cells in enumerate(rows, 2):
                    if number > MAX_ROWS + 1:
                        raise ValueError(f"最多导入 {MAX_ROWS} 行，请拆分题库")
                    values = [cell.value for cell in cells]
                    if any(value is not None for value in values):
                        if any(cell.data_type == "f" for cell in cells):
                            yield number, None, "含有公式，请粘贴为值后导入"
                        else:
                            yield number, dict(zip(header, values)), None
            finally:
                workbook.close()
        except (BadZipFile, KeyError, ParseError, SyntaxError, InvalidFileException) as error:
            raise ValueError("Excel 文件损坏，请另存为 .xlsx 或 CSV 后重试") from error
        return
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("gb18030")
        except UnicodeDecodeError as error:
            raise ValueError("无法识别文件编码，请保存为 UTF-8") from error
    if extension in (".csv", ".tsv"):
        delimiter = "\t" if extension == ".tsv" else ","
        if extension == ".csv":
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
            except csv.Error:
                pass
        reader = csv.reader(StringIO(text), delimiter=delimiter)
        header = _headers(next(reader, ()))
        for number, values in enumerate(reader, 2):
            if number > MAX_ROWS + 1:
                raise ValueError(f"最多导入 {MAX_ROWS} 行，请拆分题库")
            if not any(value.strip() for value in values):
                continue
            if len(values) != len(header):
                yield number, None, "列数与表头不一致，请检查逗号和引号"
            else:
                yield number, dict(zip(header, values)), None
    elif extension == ".jsonl":
        for number, line in enumerate(text.splitlines(), 1):
            if number > MAX_ROWS:
                raise ValueError(f"最多导入 {MAX_ROWS} 行，请拆分题库")
            if not line.strip():
                continue
            try:
                yield number, json.loads(line), None
            except ValueError:
                yield number, None, "这一行不是合法 JSON"
    else:
        try:
            value = json.loads(text)
        except ValueError as error:
            raise ValueError("JSON 文件格式有误，请检查引号、逗号和括号") from error
        if isinstance(value, dict):
            value = next((value[key] for key in ("questions", "data", "items") if isinstance(value.get(key), list)), [value])
        if not isinstance(value, list):
            raise ValueError("JSON 应为题目数组，或包含 questions / data / items 数组的对象")
        if len(value) > MAX_ROWS:
            raise ValueError(f"最多导入 {MAX_ROWS} 道题，请拆分题库")
        for number, row in enumerate(value, 1):
            yield number, row, None


def _headers(values):
    header = [str(value).strip() if value is not None else "" for value in values]
    names = [value.lower() for value in header if value]
    if not names or len(set(names)) != len(names):
        raise ValueError("第一行需要不重复的列名，例如「题目、答案、解析」")
    return header


def _identity(record):
    return compact_json({"messages": record["messages"], "tools": record.get("tools", [])})


def preview_question_bank(workspace, filename, encoded):
    if not isinstance(filename, str) or not filename or not isinstance(encoded, str):
        raise ValueError("请选择一个题库文件")
    filename = filename.replace("\\", "/").rsplit("/", 1)[-1][:180]
    extension = Path(filename).suffix.lower()
    if extension not in EXTENSIONS:
        raise ValueError("支持 .xlsx、.csv、.tsv、.json 和 .jsonl 文件")
    if len(encoded) > (MAX_UPLOAD_BYTES + 2) // 3 * 4:
        raise ValueError("题库文件不能超过 10 MiB，请拆分后导入")
    try:
        content = base64.b64decode(encoded, validate=True)
    except (ValueError, binascii.Error) as error:
        raise ValueError("文件上传不完整，请重新选择") from error
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("文件为空或超过 10 MiB")
    records, errors, seen = [], [], {}
    total = duplicates = invalid = 0
    for number, row, error in _rows(content, extension):
        total += 1
        try:
            if error:
                raise ValueError(error)
            record = _normalize(row)
            # Deduplicate the entire prompt, preventing repeated questions in both splits.
            prompt = compact_json({"messages": record["messages"][:-1], "tools": record.get("tools", [])})
            if prompt in seen:
                if seen[prompt] != _identity(record):
                    raise ValueError("相同题目的答案不一致，已保留第一次出现的答案")
                duplicates += 1
                continue
            seen[prompt] = _identity(record)
            records.append(record)
        except (ValueError, TypeError, AttributeError) as problem:
            invalid += 1
            if len(errors) < 20:
                errors.append({"row": number, "message": str(problem)})
    if not records:
        detail = errors[0]["message"] if errors else "没有发现题目"
        raise ValueError("没有可导入的题目：" + detail)
    stamp = datetime.now(timezone(timedelta(hours=8))).strftime("%Y%m%d_%H%M%S")
    identifier = "questions_" + stamp + "_" + secrets.token_hex(3)
    workspace = Path(workspace).resolve()
    package = (workspace / "data/imported" / identifier).resolve()
    if not package.is_relative_to(workspace):
        raise ValueError("题库保存位置必须位于当前项目目录内")
    package.mkdir(parents=True)
    (package / ("source" + extension)).write_bytes(content)
    (package / "records.jsonl").write_text("".join(compact_json(row) + "\n" for row in records), encoding="utf-8")
    result = {"id": identifier, "filename": filename, "total": total, "accepted": len(records),
              "duplicates": duplicates, "invalid": invalid, "errors": errors,
              "samples": records[:3], "source_sha256": sha256_file(package / ("source" + extension))}
    (package / "import.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def question_package(workspace, identifier):
    if not isinstance(identifier, str) or not re.fullmatch(r"questions_\d{8}_\d{6}_[0-9a-f]{6}", identifier):
        raise ValueError("题库编号无效，请重新选择文件")
    package = (Path(workspace) / "data/imported" / identifier).resolve()
    if not package.is_relative_to(Path(workspace).resolve()) or not (package / "import.json").is_file():
        raise ValueError("找不到上传的题库，请重新选择文件")
    return package


def prepare_question_bank(workspace, import_id, tokenizer_path=None, progress=None, cancel=None):
    workspace = Path(workspace).resolve()
    package = question_package(workspace, import_id)
    progress = progress or (lambda **value: None)

    def check():
        if cancel is not None and cancel.is_set():
            raise PreparationCancelled("已取消题库准备，上传文件保存在本机")

    check()
    records = list(iter_jsonl([package / "records.jsonl"]))
    receipt = json.loads((package / "import.json").read_text(encoding="utf-8"))
    output = workspace / "data/prepared" / import_id
    if not output.resolve().is_relative_to(workspace):
        raise ValueError("准备数据的位置必须位于当前项目目录内")
    if output.exists():
        raise ValueError("这份题库已准备过，请重新选择文件导入")
    if tokenizer_path is None:
        tokenizer_path = package / "tokenizer.json"
        if not tokenizer_path.exists():
            # Only training conversations are used to learn the new vocabulary.
            texts = [{"text": _corpus(row)} for row in records if not is_validation(_identity(row))]
            token_text = package / "tokenizer_corpus.jsonl"
            token_text.write_text("".join(compact_json(row) + "\n" for row in texts), encoding="utf-8")
            progress(phase="tokenizer", message="用题库的训练部分建立新分词器")
            train_tokenizer([token_text], tokenizer_path)
    tokenizer = LocalTokenizer(tokenizer_path)
    check()
    fit, skipped = [], 0
    for index, row in enumerate(records, 1):
        check()
        ids, _ = encode_conversation(tokenizer, row["messages"], row.get("tools"))
        if len(ids) <= 1024:
            fit.append(row)
        else:
            skipped += 1
        if index == len(records) or index % 100 == 0:
            progress(phase="filter", message="检查题目长度与训练格式", current=index, total=len(records), unit="题")
    train_count = sum(not is_validation(_identity(row)) for row in fit)
    if not train_count or train_count == len(fit):
        raise ValueError(f"有效题目不足以划分训练和验证集（可用 {len(fit)} 道，超长 {skipped} 道）。请增加不同题目，建议至少 20 道；每道题需在 1024 token 内")
    relative_tokenizer = Path(tokenizer_path).resolve().relative_to(workspace).as_posix()
    common = {"source_type": "question_bank", "package": import_id, "tokenizer_path": relative_tokenizer,
              "original_filename": receipt["filename"], "sources": [],
              "import_report": {key: receipt[key] for key in ("total", "accepted", "duplicates", "invalid", "source_sha256")}}
    common["import_report"]["too_long_skipped"] = skipped
    # Keep the same train/validation assignment across pretraining and SFT.
    corpus = {_corpus(row): _identity(row) for row in fit}
    with tempfile.TemporaryDirectory(prefix="encode-", dir=package) as temporary:
        temporary = Path(temporary)
        sft_path, corpus_path = temporary / "sft.jsonl", temporary / "corpus.jsonl"
        sft_path.write_text("".join(compact_json(row) + "\n" for row in fit), encoding="utf-8")
        corpus_path.write_text("".join(compact_json({"text": text}) + "\n" for text in corpus), encoding="utf-8")
        prepared = temporary / "prepared"
        progress(phase="encode", message="准备题库预训练数据", current=0, total=2)
        pretrain = prepare_dataset("pretrain", [corpus_path], tokenizer_path, prepared / "pretrain",
                                   metadata={**common, "display_name": receipt["filename"] + " · 题库文本"},
                                   split_keys=corpus)
        check()
        progress(phase="encode", message="准备题库问答数据", current=1, total=2)
        sft = prepare_dataset("sft", [sft_path], tokenizer_path, prepared / "sft", block_size=1023,
                             metadata={**common, "display_name": receipt["filename"] + " · 题库问答"})
        check()
        output.parent.mkdir(parents=True, exist_ok=True)
        prepared.replace(output)
    result = {"kind": "question_bank", "filename": receipt["filename"], "accepted": len(fit),
              "too_long_skipped": skipped, "duplicates": receipt["duplicates"], "invalid": receipt["invalid"],
              "prepared": {"pretrain": pretrain["splits"], "sft": sft["splits"]},
              "outputs": {"pretrain": (output / "pretrain").relative_to(workspace).as_posix(),
                          "sft": (output / "sft").relative_to(workspace).as_posix(), "tokenizer": relative_tokenizer}}
    (package / "prepared.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    progress(phase="completed", message="题库已准备好", current=2, total=2)
    return result


def _corpus(record):
    parts = []
    if record.get("tools"):
        parts.append("可用工具：" + compact_json(record["tools"]))
    for message in record["messages"]:
        content = message.get("content") or ""
        if message.get("tool_calls"):
            content += compact_json(message["tool_calls"])
        parts.append(message["role"] + "：" + content)
    return "\n".join(parts)
